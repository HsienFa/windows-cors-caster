"""Bounded, memory-only user import parsing. Never include input in errors."""

import csv
import io
import re
import struct
import zipfile
import zlib

from defusedxml import ElementTree as SafeXML
from openpyxl import Workbook, load_workbook

MAX_BYTES = 512 * 1024
MAX_ROWS = 500
MAX_ZIP_ENTRIES = 128
MAX_XML_BYTES = 8 * 1024 * 1024
HEADERS = ['username', 'password']
RULES = ('限 512 KiB；第 1 列為 username、password；最多掃描 501 列（含標題），'
         '空白列略過但仍占列數；最多 500 筆。帳號 3–50、密碼 6–100 字元，'
         '僅英文、數字、底線、連字號；去除首尾空白。XLSX 只接受文字儲存格。'
         '任一錯誤或重複即整批不匯入。')
CHARACTERS = re.compile(r'[A-Za-z0-9_-]+', re.ASCII)


class ImportProblem(Exception):
    def __init__(self, code, status=400):
        super().__init__(code)
        self.code = code
        self.status = status


def _read_zip_member(raw, archive, member, remaining):
    # ZipExtFile truncates to the declared file_size. Validate the complete raw
    # compressed stream independently, with a hard decompression output limit.
    with archive.open(member):
        pass  # Validate local header/name, encryption and overlapping entries.
    name_size, extra_size = struct.unpack_from('<HH', raw, member.header_offset + 26)
    start = member.header_offset + 30 + name_size + extra_size
    compressed = raw[start:start + member.compress_size]
    if len(compressed) != member.compress_size:
        raise ImportProblem('invalid_xlsx')
    if member.compress_type == zipfile.ZIP_DEFLATED:
        decoder = zlib.decompressobj(-15)
        content = decoder.decompress(compressed, remaining + 1)
        if len(content) > remaining:
            raise ImportProblem('xlsx_limits')
        if not decoder.eof or decoder.unused_data or decoder.unconsumed_tail:
            raise ImportProblem('invalid_xlsx')
    else:
        content = compressed
    if len(content) > remaining:
        raise ImportProblem('xlsx_limits')
    if len(content) != member.file_size or zlib.crc32(content) != member.CRC:
        raise ImportProblem('invalid_xlsx')
    return content


def _xlsx_rows(raw):
    # Validate every archive member before giving anything to openpyxl.
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_ENTRIES or sum(m.file_size for m in members) > MAX_XML_BYTES:
            raise ImportProblem('xlsx_limits')
        names = [m.filename for m in members]
        if len(set(names)) != len(names):
            raise ImportProblem('invalid_xlsx')
        total = 0
        sheets = 0
        for member in members:
            name = member.filename
            if (member.flag_bits & 1 or member.compress_type not in (0, 8)
                    or not re.fullmatch(
                        r'(\[Content_Types\]\.xml|_rels/\.rels|docProps/[A-Za-z]+\.xml|'
                        r'xl/(workbook\.xml|styles\.xml|sharedStrings\.xml|'
                        r'_rels/workbook\.xml\.rels|theme/theme[0-9]+\.xml|'
                        r'worksheets/sheet[0-9]+\.xml))', name)):
                raise ImportProblem('unsupported_xlsx')
            content = _read_zip_member(raw, archive, member, MAX_XML_BYTES - total)
            total += len(content)
            if total > MAX_XML_BYTES or len(content) != member.file_size:
                raise ImportProblem('xlsx_limits')
            root = SafeXML.fromstring(content, forbid_dtd=True, forbid_entities=True,
                                      forbid_external=True)
            worksheet = name.startswith('xl/worksheets/')
            if worksheet:
                sheets += 1
            last_row = 0
            last_column = ''
            nodes = 0
            for element in root.iter():
                nodes += 1
                if nodes > 30000:
                    raise ImportProblem('xlsx_limits')
                tag = element.tag.rsplit('}', 1)[-1]
                if tag in ('f', 'mergeCell', 'hyperlink', 'oleObject'):
                    raise ImportProblem('unsupported_xlsx')
                if tag == 'Relationship' and element.get('TargetMode') == 'External':
                    raise ImportProblem('unsupported_xlsx')
                if worksheet and tag in ('row', 'c'):
                    ref = element.get('r', '')
                    pattern = r'([AB])([1-9][0-9]*)' if tag == 'c' else r'([1-9][0-9]*)'
                    match = re.fullmatch(pattern, ref)
                    if not match or int(match.groups()[-1]) > MAX_ROWS + 1:
                        raise ImportProblem('row_or_column_limit')
                    number = int(match.groups()[-1])
                    if tag == 'row':
                        if number <= last_row:
                            raise ImportProblem('invalid_xlsx')
                        last_row, last_column = number, ''
                    else:
                        column = match.group(1)
                        if number != last_row or column <= last_column:
                            raise ImportProblem('invalid_xlsx')
                        last_column = column
            del root, content
        if sheets != 1:
            raise ImportProblem('one_sheet_required')
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False,
                             keep_links=False, keep_vba=False)
    try:
        if len(workbook.sheetnames) != 1:
            raise ImportProblem('one_sheet_required')
        sheet = workbook.active
        sheet.reset_dimensions()  # Never trust a producer's dimension declaration.
        for index, cells in enumerate(sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=2), 1):
            values = []
            for cell in cells:
                value = cell.value
                # Number format '@' is not proof of a string value.
                values.append(value if value is None or (
                    isinstance(value, str) and cell.data_type in ('s', 'inlineStr')) else False)
            yield index, values
    finally:
        workbook.close()


def _csv_rows(raw):
    text = raw.decode('utf-8-sig', errors='strict')
    reader = csv.reader(io.StringIO(text, newline=''), strict=True)
    previous = 0
    for values in reader:
        if reader.line_num > MAX_ROWS + 1:
            raise ImportProblem('row_or_column_limit')
        if reader.line_num - previous != 1:
            raise ImportProblem('multiline_field')
        previous = reader.line_num
        yield reader.line_num, values


def parse_users(raw, kind):
    """Return private tuples and safe preview rows; caller must release both."""
    if not raw or len(raw) > MAX_BYTES:
        raise ImportProblem('file_size', 413)
    if kind not in ('csv', 'xlsx'):
        raise ImportProblem('file_type')
    rows = _csv_rows(raw) if kind == 'csv' else _xlsx_rows(raw)
    private, public = [], []
    try:
        header = next(rows, None)
        if not header or header[1] != HEADERS:
            raise ImportProblem('headers')
        for row_number, values in rows:
            if not values or (len(values) <= 2 and all(
                    v is None or (isinstance(v, str) and not v.strip()) for v in values)):
                continue
            issues = []
            if len(values) != 2:
                issues.append('columns')
            username, password = (values + [None, None])[:2]
            username = username.strip() if isinstance(username, str) else ''
            password = password.strip() if isinstance(password, str) else ''
            safe_username = username if 3 <= len(username) <= 50 and CHARACTERS.fullmatch(username) else ''
            if not safe_username:
                issues.append('username')
            if not 6 <= len(password) <= 100 or not CHARACTERS.fullmatch(password):
                issues.append('password')
            public.append({'row': row_number, 'username': safe_username, 'errors': issues})
            if not issues:
                private.append((row_number, username, password))
        if not public:
            raise ImportProblem('empty_file')
        counts = {}
        for row in public:
            if row['username']:
                counts[row['username']] = counts.get(row['username'], 0) + 1
        for row in public:
            if counts.get(row['username'], 0) > 1:
                row['errors'].append('duplicate_in_file')
        return private, public
    except ImportProblem:
        raise
    except Exception:
        raise ImportProblem('invalid_file') from None
    finally:
        rows.close()


def template_bytes(kind):
    if kind == 'csv':
        return ('\ufeffusername,password\r\n').encode('utf-8')
    if kind != 'xlsx':
        raise ImportProblem('file_type')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'users'
    sheet.append(HEADERS)
    sheet.column_dimensions['A'].number_format = '@'
    sheet.column_dimensions['B'].number_format = '@'
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=2):
        for cell in row:
            cell.number_format = '@'
    # Instructions are metadata, not an extra data row or a sample password.
    workbook.properties.description = RULES
    output = io.BytesIO()
    try:
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
        output.close()
