"""Synthetic files and per-test SQLite only. No production config, logs or listeners."""
import io
import secrets
import sqlite3
import struct
import threading
import time
import unittest
import zipfile
import zlib
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

import test_rover_web_api as fixture  # Establish isolated config and logger before src imports.
from src import database, user_import as parser, user_import_web
from openpyxl import Workbook, load_workbook
from defusedxml.common import DefusedXmlException


def xlsx(rows, mutate=None):
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    if mutate:
        mutate(book)
    data = io.BytesIO()
    book.save(data)
    book.close()
    return data.getvalue()


class ParserTests(unittest.TestCase):
    def test_forged_zip_size_cannot_hide_decompressed_bytes(self):
        output = io.BytesIO()
        source = xlsx([parser.HEADERS, ['user001', 'abcdef']])
        with zipfile.ZipFile(io.BytesIO(source)) as original, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
            for name in original.namelist():
                target.writestr(name, original.read(name))
            target.writestr('xl/theme/theme99.xml', b'<x/>' + b' ' * 32)
        raw = bytearray(output.getvalue())
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entry = archive.getinfo('xl/theme/theme99.xml')
        struct.pack_into('<I', raw, entry.header_offset + 14, zlib.crc32(b'<x/>'))
        struct.pack_into('<I', raw, entry.header_offset + 22, 4)
        position = raw.find(b'PK\x01\x02')
        while raw[position:position + 4] == b'PK\x01\x02':
            name_len, extra_len, comment_len = struct.unpack_from('<HHH', raw, position + 28)
            if raw[position + 46:position + 46 + name_len] == b'xl/theme/theme99.xml':
                struct.pack_into('<I', raw, position + 16, zlib.crc32(b'<x/>'))
                struct.pack_into('<I', raw, position + 24, 4)
                break
            position += 46 + name_len + extra_len + comment_len
        else:
            self.fail('synthetic central entry missing')
        with mock.patch.object(parser, 'load_workbook') as load, self.assertRaises(parser.ImportProblem):
            parser.parse_users(bytes(raw), 'xlsx')
        load.assert_not_called()

    def test_csv_bom_and_strip_preserve_leading_zero(self):
        for bom in ('', '\ufeff'):
            private, public = parser.parse_users((bom + 'username,password\n 001 , 001234 \n').encode(), 'csv')
            self.assertEqual(private, [(2, '001', '001234')])
            self.assertEqual(public, [{'row': 2, 'username': '001', 'errors': []}])

    def test_rows_and_blank_lines_have_same_limits_for_both_formats(self):
        for kind in ('csv', 'xlsx'):
            def encode(rows):
                return xlsx(rows) if kind == 'xlsx' else ('\n'.join(','.join(r) for r in rows) + '\n').encode()
            rows = [parser.HEADERS] + [['', '']] * 499 + [['user500', 'secret500']]
            private, public = parser.parse_users(encode(rows), kind)
            self.assertEqual(public[0]['row'], 501)
            self.assertEqual(len(private), 1)
            with self.assertRaises(parser.ImportProblem):
                parser.parse_users(encode(rows + [['late', 'secret501']]), kind)

    def test_template_only_headers_and_text_format(self):
        self.assertEqual(parser.template_bytes('csv'), b'\xef\xbb\xbfusername,password\r\n')
        book = load_workbook(io.BytesIO(parser.template_bytes('xlsx')))
        self.addCleanup(book.close)
        self.assertEqual([c.value for c in book.active[1]], parser.HEADERS)
        self.assertEqual(book.active['B2'].number_format, '@')
        self.assertTrue(all(c.value is None for row in book.active.iter_rows(min_row=2) for c in row))
        self.assertIn('501', book.properties.description)

    def test_numeric_password_even_text_formatted_is_rejected(self):
        raw = xlsx([parser.HEADERS, ['testuser', 123456]], lambda b: setattr(b.active['B2'], 'number_format', '@'))
        _, public = parser.parse_users(raw, 'xlsx')
        self.assertEqual(public[0]['errors'], ['password'])

    def test_dates_formula_merge_extra_sheet_and_links_rejected(self):
        from datetime import datetime
        mutations = [lambda b: setattr(b.active['B2'], 'value', '=1+1'),
                     lambda b: b.active.merge_cells('A2:B2'), lambda b: b.create_sheet('extra'),
                     lambda b: setattr(b.active['A2'], 'hyperlink', 'https://example.invalid')]
        for mutate in mutations:
            with self.subTest(mutation=mutations.index(mutate)), self.assertRaises(parser.ImportProblem):
                parser.parse_users(xlsx([parser.HEADERS, ['testuser', 'abcdef']], mutate), 'xlsx')
        _, rows = parser.parse_users(xlsx([parser.HEADERS, ['testuser', datetime(2020, 1, 1)]]), 'xlsx')
        self.assertIn('password', rows[0]['errors'])

    def test_header_multiline_empty_encoding_and_size(self):
        for raw in (b'username,password,extra\n', b'password,username\n', b'username,username\n',
                    b'username,password\n"bad\nname",abcdef\n', b'username,password\n', b'\xff\xfe'):
            with self.subTest(raw_length=len(raw)), self.assertRaises(parser.ImportProblem):
                parser.parse_users(raw, 'csv')
        with self.assertRaises(parser.ImportProblem):
            parser.parse_users(b'x' * (parser.MAX_BYTES + 1), 'csv')

    def test_duplicate_all_rows_and_no_password_in_preview(self):
        _, rows = parser.parse_users(b'username,password\nabc,secretA\nabc,secretB\nABC,secretC\nbad<,secretD\n', 'csv')
        self.assertIn('duplicate_in_file', rows[0]['errors'])
        self.assertIn('duplicate_in_file', rows[1]['errors'])
        self.assertEqual(rows[2]['errors'], [])
        self.assertEqual(rows[3]['username'], '')
        self.assertNotIn('secret', str(rows))

    def test_zip_limits_and_xml_defenses_run_before_openpyxl(self):
        for count, data in ((129, b'<x/>'), (1, b' ' * (parser.MAX_XML_BYTES + 1)),
                            (1, b'<!DOCTYPE x [<!ENTITY a "secret">]><x>&a;</x>')):
            output = io.BytesIO()
            with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
                for n in range(count):
                    archive.writestr(f'xl/theme/theme{n}.xml', data)
            with mock.patch.object(parser, 'load_workbook') as load, self.assertRaises(parser.ImportProblem):
                parser.parse_users(output.getvalue(), 'xlsx')
            load.assert_not_called()
        from openpyxl.xml.functions import iterparse
        from defusedxml.ElementTree import iterparse as protected_iterparse
        self.assertIs(iterparse, protected_iterparse)
        with self.assertRaises(DefusedXmlException):
            list(iterparse(io.BytesIO(b'<!DOCTYPE x [<!ENTITY a "secret">]><x>&a;</x>')))

    def test_dimension_does_not_hide_late_rows(self):
        raw = xlsx([parser.HEADERS, ['user001', 'secretA']])
        result = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(result, 'w') as target:
            for name in source.namelist():
                data = source.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'r="2"', b'r="502"').replace(b'A2', b'A502').replace(b'B2', b'B502')
                target.writestr(name, data)
        with self.assertRaises(parser.ImportProblem):
            parser.parse_users(result.getvalue(), 'xlsx')

    def test_extra_blank_columns_and_duplicate_xlsx_rows_rejected(self):
        _, rows = parser.parse_users(b'username,password\nuser001,abcdef\n,,\n', 'csv')
        self.assertIn('columns', rows[1]['errors'])
        raw = xlsx([parser.HEADERS, ['user001', 'abcdef']])
        for needle, replacement in ((b'r="2"', b'r="1"'), (b'r="B2"', b'r="A2"')):
            output = io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, 'w') as target:
                for name in source.namelist():
                    data = source.read(name)
                    if name == 'xl/worksheets/sheet1.xml':
                        data = data.replace(needle, replacement)
                    target.writestr(name, data)
            with mock.patch.object(parser, 'load_workbook') as load, self.assertRaises(parser.ImportProblem):
                parser.parse_users(output.getvalue(), 'xlsx')
            load.assert_not_called()


class ImportAPITests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = str(Path(self.temp.name) / 'synthetic.db')
        self.patch = mock.patch.object(database.config, 'DATABASE_PATH', self.path)
        self.patch.start()
        self.addCleanup(self.patch.stop)
        connections = []
        real_connect = sqlite3.connect
        class TrackedConnection(sqlite3.Connection):
            closed = False
            def close(self):
                super().close()
                self.closed = True
        def connect(*args, **kwargs):
            kwargs['factory'] = TrackedConnection
            conn = real_connect(*args, **kwargs)
            connections.append(conn)
            return conn
        connection_patch = mock.patch.object(sqlite3, 'connect', side_effect=connect)
        connection_patch.start()
        self.addCleanup(connection_patch.stop)
        # Some existing read-only DB methods do not close their context-manager
        # connections; explicitly close all synthetic handles before temp cleanup.
        self.addCleanup(lambda: [conn.close() for conn in connections if not conn.closed])
        with sqlite3.connect(self.path) as conn:
            conn.executescript('CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE, password TEXT NOT NULL);'
                               'CREATE TABLE mounts (id INTEGER PRIMARY KEY, mount TEXT, password TEXT, user_id INTEGER);'
                               "INSERT INTO mounts VALUES (1, 'SYNTHETIC', '', NULL);")
        self.db = database.DatabaseManager()
        self.db.verify_admin = mock.Mock(return_value=True)
        self.web = fixture.web.WebManager(self.db, mock.Mock(), 0)
        self.web.app.secret_key = secrets.token_hex(32)
        self.web.app.config['TESTING'] = True
        self.http, self.headers = self.login()
        self.raw = b'username,password\nuser001,SyntheticPassword\nuser002,SyntheticPassword\n'

    def login(self):
        http = self.web.app.test_client()
        self.assertEqual(http.post('/api/login', json={'username': 'testadmin', 'password': secrets.token_hex(12)}).status_code, 200)
        headers = {'Origin': 'http://localhost', 'X-Import-Request': '1', 'Content-Type': 'application/octet-stream', 'X-Import-Format': 'csv'}
        response = http.post('/api/users/import/context', headers=headers)
        self.assertEqual(response.status_code, 200)
        headers['X-Import-CSRF'] = response.get_json()['csrf']
        return http, headers

    def post(self, action, raw=None, headers=None, http=None):
        return (http or self.http).post('/api/users/import/' + action, data=self.raw if raw is None else raw,
                                       headers=headers or self.headers)

    def preview(self):
        response = self.post('preview')
        self.assertEqual(response.status_code, 200)
        self.headers['X-Import-Preview'] = response.get_json()['preview']
        return response

    def names(self):
        with sqlite3.connect(self.path) as conn:
            return [r[0] for r in conn.execute('SELECT username FROM users ORDER BY username')]

    def test_preview_no_write_confirm_reuses_hash_and_ntrip_verification(self):
        response = self.preview()
        self.assertEqual(self.names(), [])
        self.assertNotIn('SyntheticPassword', response.get_data(as_text=True))
        self.assertNotIn('SyntheticPassword', str(self.web.user_import.previews))
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertEqual(self.post('confirm').status_code, 201)
        with sqlite3.connect(self.path) as conn:
            hashes = [r[0] for r in conn.execute('SELECT password FROM users')]
        self.assertTrue(all(database.verify_password(h, 'SyntheticPassword') for h in hashes))
        self.assertNotEqual(hashes[0], hashes[1])
        self.assertTrue(self.db.verify_download_user('SYNTHETIC', 'user001', 'SyntheticPassword')[0])
        self.assertFalse(self.db.verify_download_user('SYNTHETIC', 'user001', 'wrong')[0])
        self.assertEqual(self.post('confirm').status_code, 409)
        self.assertEqual(self.web.user_import.previews, {})

    def test_preview_expires_during_hash_rolls_back(self):
        self.preview()
        deadline = next(iter(self.web.user_import.previews.values()))['deadline']
        real_time = time.monotonic
        expired = False
        def clock():
            return deadline + 1 if expired else real_time()
        def hashing(password):
            nonlocal expired
            expired = True
            return database.hash_password(password)
        with mock.patch.object(user_import_web.time, 'monotonic', side_effect=clock), \
                mock.patch.object(user_import_web, 'hash_password', side_effect=hashing):
            response = self.post('confirm')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json(), {'error': 'preview_expired'})
        self.assertEqual(self.names(), [])
        self.assertEqual(self.web.user_import.previews, {})

    def test_anonymous_csrf_and_origin_all_routes(self):
        anonymous = self.web.app.test_client()
        for action in ('context', 'preview', 'confirm', 'cancel', 'template-csv', 'template-xlsx'):
            self.assertEqual(self.post(action, http=anonymous).status_code, 401)
            headers = dict(self.headers, Origin='https://example.invalid')
            self.assertEqual(self.post(action, headers=headers).status_code, 403)
            if action != 'context':
                headers = dict(self.headers, **{'X-Import-CSRF': 'invalid'})
                self.assertEqual(self.post(action, headers=headers).status_code, 403)

    def test_forbidden_request_does_not_cancel_valid_preview(self):
        self.preview()
        for headers in (dict(self.headers, Origin='https://example.invalid'),
                        dict(self.headers, **{'X-Import-CSRF': 'invalid'})):
            self.assertEqual(self.post('cancel', headers=headers).status_code, 403)
            self.assertEqual(len(self.web.user_import.previews), 1)
        self.assertEqual(self.post('confirm').status_code, 201)

    def test_stale_confirmation_does_not_consume_newer_preview(self):
        self.preview()
        old_headers = dict(self.headers)
        self.preview()
        self.assertEqual(self.post('confirm', headers=old_headers).status_code, 409)
        self.assertEqual(len(self.web.user_import.previews), 1)
        self.assertEqual(self.post('confirm').status_code, 201)

    def test_older_parse_failure_does_not_remove_newer_preview(self):
        entered, release = threading.Event(), threading.Event()
        original = user_import_web.parse_users
        old_raw = b'username,password\nolduser,abcdef\n'
        def parse(raw, kind):
            if raw == old_raw:
                entered.set()
                if not release.wait(3):
                    raise AssertionError('bounded wait')
                raise parser.ImportProblem('invalid_file')
            return original(raw, kind)
        results = []
        worker = threading.Thread(target=lambda: results.append(self.post('preview', old_raw)), daemon=True)
        with mock.patch.object(user_import_web, 'parse_users', side_effect=parse):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                self.preview()
            finally:
                release.set()
                worker.join(3)
        self.assertFalse(worker.is_alive())
        self.assertEqual(results[0].status_code, 400)
        self.assertEqual(len(self.web.user_import.previews), 1)
        self.assertEqual(self.post('confirm').status_code, 201)

    def test_auth_lock_wait_rolls_back_and_releases_database_lock(self):
        self.preview()
        entered, release, done = threading.Event(), threading.Event(), threading.Event()
        original = user_import_web.hash_password
        def hashing(password):
            entered.set()
            if not release.wait(4):
                raise AssertionError('bounded wait')
            return original(password)
        results = []
        def confirm():
            results.append(self.post('confirm'))
            done.set()
        worker = threading.Thread(target=confirm, daemon=True)
        with mock.patch.object(user_import_web, 'hash_password', side_effect=hashing):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                with self.web._admin_lock:
                    release.set()
                    completed_while_locked = done.wait(3)
            finally:
                release.set()
                worker.join(3)
        self.assertFalse(worker.is_alive())
        self.assertTrue(completed_while_locked)
        self.assertEqual(results[0].status_code, 503)
        self.assertFalse(database.db_lock.locked())
        self.assertEqual(self.names(), [])
        self.assertEqual(self.web.user_import.previews, {})

    def test_concurrent_single_add_during_confirm_causes_full_rollback(self):
        self.preview()
        other, _ = self.login()
        entered, release = threading.Event(), threading.Event()
        original = user_import_web.hash_password
        def hashing(password):
            entered.set()
            if not release.wait(3):
                raise AssertionError('bounded wait')
            return original(password)
        results = []
        worker = threading.Thread(target=lambda: results.append(self.post('confirm')), daemon=True)
        with mock.patch.object(user_import_web, 'hash_password', side_effect=hashing):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                added = other.post('/api/users', json={'username': 'user002', 'password': 'OriginalPassword'})
                self.assertEqual(added.status_code, 201)
            finally:
                release.set()
                worker.join(3)
        self.assertFalse(worker.is_alive())
        self.assertEqual(results[0].status_code, 409)
        self.assertEqual(self.names(), ['user002'])
        self.assertTrue(database.verify_password(self.db.get_user_password('user002'), 'OriginalPassword'))

    def test_response_failure_after_commit_does_not_claim_rollback(self):
        self.preview()
        original = self.web.user_import.response
        def response(payload, status=200):
            if status == 201:
                raise RuntimeError('SyntheticPassword')
            return original(payload, status)
        with mock.patch.object(self.web.user_import, 'response', side_effect=response):
            result = self.post('confirm')
        self.assertEqual(result.status_code, 503)
        self.assertEqual(result.get_json(), {'error': 'commit_result_unknown'})
        self.assertNotIn('SyntheticPassword', result.get_data(as_text=True))
        self.assertEqual(self.names(), ['user001', 'user002'])
        self.assertEqual(self.post('confirm').status_code, 409)

    def test_preview_response_failure_does_not_keep_unusable_record(self):
        original = self.web.user_import.response
        def response(payload, status=200):
            if 'preview' in payload:
                raise RuntimeError('SyntheticPassword')
            return original(payload, status)
        with mock.patch.object(self.web.user_import, 'response', side_effect=response):
            result = self.post('preview')
        self.assertEqual(result.status_code, 400)
        self.assertEqual(self.web.user_import.previews, {})
        self.assertEqual(self.names(), [])

    def test_close_failure_after_commit_releases_lock_and_reports_uncertainty(self):
        self.preview()
        original = sqlite3.connect
        class ClosingFailure:
            def __init__(self, conn):
                self.conn = conn
            def __getattr__(self, name):
                return getattr(self.conn, name)
            def close(self):
                self.conn.close()
                raise RuntimeError('SyntheticPassword')
        with mock.patch.object(self.db, 'existing_import_users', return_value=set()), \
                mock.patch.object(sqlite3, 'connect', side_effect=lambda *a, **kw: ClosingFailure(original(*a, **kw))):
            result = self.post('confirm')
        self.assertEqual(result.status_code, 503)
        self.assertEqual(result.get_json(), {'error': 'commit_result_unknown'})
        self.assertFalse(database.db_lock.locked())
        self.assertEqual(self.names(), ['user001', 'user002'])

    def test_download_templates_authenticated_no_store(self):
        for kind in ('csv', 'xlsx'):
            response = self.post('template-' + kind)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers['Cache-Control'], 'no-store')
            self.assertIn('attachment', response.headers['Content-Disposition'])

    def test_invalid_rows_duplicate_and_existing_entire_batch_rejected(self):
        for raw in (b'username,password\nuser001,abcdef\nuser001,ghijkl\n',
                    b'username,password\nuser001,abcdef\nx,abcdef\n'):
            response = self.post('preview', raw)
            self.assertEqual(response.status_code, 400)
            self.assertFalse(response.get_json()['can_import'])
            self.assertEqual(self.names(), [])
        self.db.add_user('user001', 'OriginalPassword')
        self.assertEqual(self.post('preview').status_code, 400)
        self.assertEqual(self.names(), ['user001'])
        self.assertTrue(database.verify_password(self.db.get_user_password('user001'), 'OriginalPassword'))

    def test_preview_then_single_insert_conflict_and_replay(self):
        self.preview()
        self.db.add_user('user002', 'OriginalPassword')
        self.assertEqual(self.post('confirm').status_code, 400)
        self.assertEqual(self.names(), ['user002'])
        self.assertEqual(self.post('confirm').status_code, 409)

    def test_cancel_changed_file_and_expiry_cleanup(self):
        for mode in ('cancel', 'changed', 'expired'):
            self.preview()
            if mode == 'cancel':
                self.post('cancel')
            elif mode == 'expired':
                for record in self.web.user_import.previews.values():
                    record['deadline'] = time.monotonic() - 1
            raw = self.raw + b'user003,abcdef\n' if mode == 'changed' else self.raw
            self.assertEqual(self.post('confirm', raw).status_code, 409)
            self.assertEqual(self.names(), [])
            self.assertEqual(self.web.user_import.previews, {})

    def test_logout_old_cookie_and_independent_login(self):
        self.preview()
        other, headers = self.login()
        headers['X-Import-Preview'] = self.headers['X-Import-Preview']
        self.assertEqual(self.post('confirm', http=other, headers=headers).status_code, 409)
        old = self.http.get_cookie('session').value
        self.http.post('/logout')
        self.assertEqual(self.web.user_import.previews, {})
        self.http.set_cookie('session', old)
        self.assertEqual(self.post('confirm').status_code, 401)
        self.assertEqual(self.post('preview', http=other, headers=headers).status_code, 200)

    def test_login_expiry_denies_even_existing_preview(self):
        self.preview()
        for grant in self.web._admin_logins.values():
            grant['expires_at'] = time.monotonic() - 1
        self.assertEqual(self.post('confirm').status_code, 401)
        self.assertEqual(self.web.user_import.previews, {})

    def test_actual_read_bounded_without_content_length(self):
        class LimitedRead(io.BytesIO):
            sizes = []
            def readinto(self, buffer):
                self.sizes.append(len(buffer))
                return super().readinto(buffer)
        stream = LimitedRead(b'x' * (parser.MAX_BYTES + 100))
        response = self.http.open('/api/users/import/preview', method='POST', headers=self.headers,
                                 environ_overrides={'wsgi.input': stream, 'wsgi.input_terminated': True,
                                                    'CONTENT_LENGTH': ''})
        self.assertEqual(response.status_code, 413)
        self.assertLessEqual(stream.tell(), parser.MAX_BYTES + 1)

    def test_concurrency_and_preview_capacity_bounds(self):
        self.web.user_import.slots.acquire()
        self.web.user_import.slots.acquire()
        try:
            self.assertEqual(self.post('preview').status_code, 503)
        finally:
            self.web.user_import.slots.release()
            self.web.user_import.slots.release()
        for _ in range(32):
            client, headers = self.login()
            self.assertEqual(self.post('preview', http=client, headers=headers).status_code, 200)
        self.assertEqual(self.post('preview').status_code, 503)
        self.assertEqual(len(self.web.user_import.previews), 32)

    def test_mid_insert_failure_rolls_back_all_and_hides_exception(self):
        self.preview()
        with sqlite3.connect(self.path) as conn:
            conn.execute("CREATE TRIGGER fail_second BEFORE INSERT ON users WHEN NEW.username='user002' BEGIN SELECT RAISE(ABORT, 'SyntheticPassword'); END;")
        response = self.post('confirm')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.names(), [])
        self.assertNotIn('SyntheticPassword', response.get_data(as_text=True))

    def test_hashing_has_no_db_or_auth_lock_and_logout_during_hash_denies(self):
        self.preview()
        entered, release = threading.Event(), threading.Event()
        original = user_import_web.hash_password
        def hashing(value):
            self.assertFalse(self.web._admin_lock._is_owned())
            self.assertTrue(database.db_lock.acquire(blocking=False))
            database.db_lock.release()
            entered.set()
            if not release.wait(3):
                raise AssertionError('bounded wait')
            return original(value)
        results = []
        worker = threading.Thread(target=lambda: results.append(self.post('confirm')), daemon=True)
        with mock.patch.object(user_import_web, 'hash_password', side_effect=hashing):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                self.http.post('/logout')
            finally:
                release.set()
                worker.join(3)
        self.assertFalse(worker.is_alive())
        self.assertEqual(results[0].status_code, 401)
        self.assertEqual(self.names(), [])

    def test_cancel_during_parse_prevents_late_preview(self):
        entered, release = threading.Event(), threading.Event()
        original = user_import_web.parse_users
        def parse(*args):
            entered.set()
            if not release.wait(3):
                raise AssertionError('bounded wait')
            return original(*args)
        results = []
        worker = threading.Thread(target=lambda: results.append(self.post('preview')), daemon=True)
        with mock.patch.object(user_import_web, 'parse_users', side_effect=parse):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                self.post('cancel')
            finally:
                release.set()
                worker.join(3)
        self.assertFalse(worker.is_alive())
        self.assertEqual(results[0].status_code, 409)
        self.assertEqual(self.web.user_import.previews, {})

    def test_transaction_rechecks_conflict_after_preflight(self):
        self.preview()
        original = self.db.import_users_atomic
        def concurrent_insert(rows, guard):
            self.db.add_user('user002', 'OriginalPassword')
            return original(rows, guard)
        with mock.patch.object(self.db, 'import_users_atomic', side_effect=concurrent_insert):
            self.assertEqual(self.post('confirm').status_code, 409)
        self.assertEqual(self.names(), ['user002'])

    def test_single_add_two_character_compatibility_unchanged(self):
        response = self.http.post('/api/users', json={'username': 'ab', 'password': 'SyntheticPassword'})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(self.names(), ['ab'])

    def test_database_lock_wait_is_bounded(self):
        self.preview()
        database.db_lock.acquire()
        try:
            started = time.monotonic()
            response = self.post('confirm')
            self.assertEqual(response.status_code, 503)
            self.assertLess(time.monotonic() - started, 4)
        finally:
            database.db_lock.release()
        self.assertEqual(self.names(), [])

    def test_500_records_fit_and_501_rejected(self):
        raw = ('username,password\n' + ''.join(f'user{i:03},SyntheticPassword\n' for i in range(500))).encode()
        preview = self.post('preview', raw)
        self.assertEqual(preview.get_json()['valid_count'], 500)
        self.headers['X-Import-Preview'] = preview.get_json()['preview']
        self.assertEqual(self.post('confirm', raw).status_code, 201)
        self.assertEqual(len(self.names()), 500)
        self.assertEqual(self.post('preview', raw + b'lastUser,SyntheticPassword\n').status_code, 400)

    def test_duplicate_confirmation_during_hash_inserts_once(self):
        self.preview()
        entered, release = threading.Event(), threading.Event()
        original = user_import_web.hash_password
        def hashing(value):
            entered.set()
            if not release.wait(3):
                raise AssertionError('bounded wait')
            return original(value)
        results = []
        worker = threading.Thread(target=lambda: results.append(self.post('confirm')), daemon=True)
        with mock.patch.object(user_import_web, 'hash_password', side_effect=hashing):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                self.assertEqual(self.post('confirm').status_code, 409)
            finally:
                release.set()
                worker.join(3)
        self.assertFalse(worker.is_alive())
        self.assertEqual(results[0].status_code, 201)
        self.assertEqual(self.names(), ['user001', 'user002'])

    def test_commit_before_logout_has_definite_order(self):
        self.preview()
        entered, release, done = threading.Event(), threading.Event(), threading.Event()
        original = self.web.user_import.commit_guard
        @contextmanager
        def guard(epoch=None, deadline=None):
            with original(epoch, deadline):
                # Only hold at the transaction's final guard, not the preflight guard.
                if database.db_lock.locked():
                    entered.set()
                    if not release.wait(3):
                        raise AssertionError('bounded wait')
                yield
        results = []
        worker = threading.Thread(target=lambda: results.append(self.post('confirm')), daemon=True)
        logout = threading.Thread(target=lambda: (self.http.post('/logout'), done.set()), daemon=True)
        with mock.patch.object(self.web.user_import, 'commit_guard', side_effect=guard):
            worker.start()
            try:
                self.assertTrue(entered.wait(3))
                logout.start()
                self.assertFalse(done.wait(.05))
            finally:
                release.set()
                worker.join(3)
                if logout.ident:
                    logout.join(3)
        self.assertFalse(worker.is_alive())
        self.assertTrue(done.is_set())
        self.assertEqual(results[0].status_code, 201)
        self.assertEqual(self.names(), ['user001', 'user002'])


if __name__ == '__main__':
    unittest.main()
